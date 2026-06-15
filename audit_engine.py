#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
응용시스템 감리 자동 점검 엔진
================================
NIA 감리점검해설서 V3.0(SD/응용시스템) 기반 단계별 산출물 자동 점검.

rules/audit_rules.yaml 룰셋을 읽어 산출물 존재/내용/금지토큰/추적성을 검증하고,
- 콘솔 요약
- audit_report.md (사람용)
- audit_report.json (CI 연동용)
을 생성한다. 미충족 항목이 있으면 종료코드 1을 반환 → CI에서 머지/제출 차단 가능.

사용법:
    python audit_engine.py --rules rules/audit_rules.yaml [--phase SD22-2] [--strict]

의존성: PyYAML  (텍스트 추출용 선택: python-docx, pdfplumber, openpyxl)
"""
import argparse
import fnmatch
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import yaml
except ImportError:
    sys.exit("PyYAML이 필요합니다:  pip install pyyaml")


# ----------------------------- 텍스트 추출 -----------------------------
def extract_text(path: Path) -> str:
    """산출물 파일에서 텍스트를 추출. 미지원 포맷은 빈 문자열."""
    suffix = path.suffix.lower()
    try:
        if suffix in (".md", ".txt", ".csv"):
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".docx":
            import docx  # python-docx
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if suffix == ".pdf":
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                return "\n".join((pg.extract_text() or "") for pg in pdf.pages)
        if suffix == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    out.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(out)
    except Exception as e:  # 추출 실패해도 점검은 계속 (존재 점검은 통과)
        return f"__EXTRACT_ERROR__:{e}"
    return ""  # hwp 등 미지원


# ----------------------------- 파일 매칭 -----------------------------
def expand_braces(pattern: str):
    """foo.{md,docx} → [foo.md, foo.docx]"""
    m = re.search(r"\{([^}]+)\}", pattern)
    if not m:
        return [pattern]
    options = m.group(1).split(",")
    return [pattern[: m.start()] + opt + pattern[m.end():] for opt in options]


def find_files(doc_root: Path, pattern: str):
    matches = []
    for pat in expand_braces(pattern):
        # ** 재귀 지원
        for p in doc_root.glob(pat):
            if p.is_file():
                matches.append(p)
    return sorted(set(matches))


# ----------------------------- 점검 로직 -----------------------------
class Result:
    def __init__(self, code, desc, status, detail=""):
        self.code, self.desc, self.status, self.detail = code, desc, status, detail

    def to_dict(self):
        return {"code": self.code, "desc": self.desc,
                "status": self.status, "detail": self.detail}


def run_check(check, deliv_files, deliv_text, common, rtm_path):
    code = check.get("code", "?")
    desc = check.get("desc", "")
    ctype = check["type"]
    target = check.get("target")
    files = deliv_files.get(target, [])
    text = deliv_text.get(target, "")

    # 대상 산출물이 아예 없으면 우선 결함
    if ctype != "traceability" and target and not files:
        return Result(code, desc, "FAIL", f"대상 산출물({target}) 없음")

    if ctype == "contains":
        missing = [k for k in check["keywords"] if k not in text]
        if missing:
            return Result(code, desc, "FAIL", f"누락 키워드: {', '.join(missing)}")
        return Result(code, desc, "PASS", f"{len(check['keywords'])}개 키워드 확인")

    if ctype == "not_contains":
        found = [k for k in check["tokens"] if k in text]
        if found:
            return Result(code, desc, "FAIL", f"금지 토큰 발견: {', '.join(found)}")
        return Result(code, desc, "PASS")

    if ctype == "min_count":
        n = len(re.findall(check["pattern"], text))
        if n < check["min"]:
            return Result(code, desc, "FAIL", f"{check['pattern']} {n}건 < 최소 {check['min']}건")
        return Result(code, desc, "PASS", f"{n}건")

    return Result(code, desc, "SKIP", f"미지원 타입: {ctype}")


def check_traceability(rtm_path: Path):
    """RTM.csv 추적성 점검. 헤더: REQ_ID,UC_ID,CLS_ID,CMP_ID,UT_ID,UAT_ID"""
    results = []
    if not rtm_path.exists():
        return [Result("RTM", "요구사항 추적 매트릭스", "WARN", "RTM.csv 없음")]
    import csv
    chain = ["UC_ID", "CLS_ID", "CMP_ID", "UT_ID", "UAT_ID"]
    orphans = []
    with rtm_path.open(encoding="utf-8") as f:
        for row in csv.DictReader(f):
            req = row.get("REQ_ID", "").strip()
            if not req:
                continue
            gaps = [c for c in chain if not row.get(c, "").strip()]
            if gaps:
                orphans.append(f"{req}(단절:{','.join(gaps)})")
    if orphans:
        results.append(Result("RTM-TRACE", "요구사항 추적성 단절",
                              "FAIL", f"{len(orphans)}건: " + "; ".join(orphans[:10])))
    else:
        results.append(Result("RTM-TRACE", "요구사항 추적성", "PASS", "전 요구사항 추적 완결"))
    return results


def check_common_format(deliv_text, common):
    """전 산출물 공통 — 금지 토큰 점검."""
    results = []
    forbidden = common.get("forbidden_tokens", [])
    for dname, text in deliv_text.items():
        found = [t for t in forbidden if t in text]
        if found:
            results.append(Result(f"FMT-{dname}", f"{dname} 금지토큰",
                                  "FAIL", f"발견: {', '.join(found)}"))
    if not results:
        results.append(Result("FMT-COMMON", "공통 형식(금지토큰)", "PASS"))
    return results


# ----------------------------- 메인 -----------------------------
def main():
    ap = argparse.ArgumentParser(description="응용시스템 감리 자동 점검 엔진")
    ap.add_argument("--rules", default="rules/audit_rules.yaml")
    ap.add_argument("--phase", help="특정 단계만 점검 (예: SD22-2)")
    ap.add_argument("--strict", action="store_true", help="WARN도 실패로 간주")
    args = ap.parse_args()

    rules = yaml.safe_load(Path(args.rules).read_text(encoding="utf-8"))
    base = Path(args.rules).resolve().parent.parent
    doc_root = (base / rules["project"]["doc_root"]).resolve()
    rtm_path = (base / rules["project"]["rtm_file"]).resolve()
    common = rules.get("common_checks", {})

    all_results = []
    phase_summary = []

    for phase in rules["phases"]:
        if args.phase and phase["code"] != args.phase:
            continue
        pcode, pname = phase["code"], phase["name"]
        deliv_files, deliv_text = {}, {}

        # 1) 산출물 존재 점검
        for d in phase.get("deliverables", []):
            files = find_files(doc_root, d["pattern"])
            deliv_files[d["id"]] = files
            txt = "\n".join(extract_text(f) for f in files)
            deliv_text[d["id"]] = txt
            status = "PASS" if files else "FAIL"
            all_results.append(Result(
                f"{pcode}/{d['id']}", f"[산출물] {d['name']}",
                status, f"{len(files)}개 파일" if files else "산출물 없음"))

        # 2) 점검 룰 실행
        for chk in phase.get("checks", []):
            all_results.append(run_check(chk, deliv_files, deliv_text, common, rtm_path))

        # 3) 공통 형식 점검 (단계 산출물 대상)
        all_results.extend(check_common_format(deliv_text, common))

        p_results = [r for r in all_results if r.code.startswith(pcode) or r.code.startswith("FMT")]
        fails = sum(1 for r in all_results if r.status == "FAIL"
                    and (r.code.startswith(pcode)))
        phase_summary.append((pcode, pname, fails))

    # 4) 추적성 점검 (전체)
    if not args.phase:
        all_results.extend(check_traceability(rtm_path))

    # ----------------------------- 리포트 -----------------------------
    fail = sum(1 for r in all_results if r.status == "FAIL")
    warn = sum(1 for r in all_results if r.status == "WARN")
    pas = sum(1 for r in all_results if r.status == "PASS")
    total = len(all_results)

    icon = {"PASS": "✅", "FAIL": "❌", "WARN": "⚠️", "SKIP": "⏭️"}
    print("=" * 64)
    print(f"  감리 자동 점검 결과   (PASS {pas} / FAIL {fail} / WARN {warn} / 총 {total})")
    print("=" * 64)
    for r in all_results:
        print(f"  {icon.get(r.status,'?')} [{r.code}] {r.desc}"
              + (f"  — {r.detail}" if r.detail else ""))

    # MD 리포트
    md = [f"# 감리 자동 점검 리포트", "",
          f"- 생성: {datetime.now():%Y-%m-%d %H:%M}",
          f"- 결과: **PASS {pas} / FAIL {fail} / WARN {warn}** (총 {total})", "",
          "| 상태 | 코드 | 점검 내용 | 상세 |", "|------|------|-----------|------|"]
    for r in all_results:
        md.append(f"| {icon.get(r.status,'?')} {r.status} | `{r.code}` | {r.desc} | {r.detail} |")
    Path(base / "audit_report.md").write_text("\n".join(md), encoding="utf-8")

    # JSON 리포트 (CI 연동)
    Path(base / "audit_report.json").write_text(
        json.dumps({"summary": {"pass": pas, "fail": fail, "warn": warn, "total": total},
                    "results": [r.to_dict() for r in all_results]},
                   ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n  리포트: audit_report.md / audit_report.json")
    if fail > 0 or (args.strict and warn > 0):
        print("  ❌ 감리 제출 기준 미충족 — 결함을 해소하세요.")
        sys.exit(1)
    print("  ✅ 감리 제출 기준 충족.")
    sys.exit(0)


if __name__ == "__main__":
    main()

